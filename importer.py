"""
.txt, .csv va .docx fayllardan (english, uzbek, example) ma'lumotlarini aqlli tarzda ajratib oladi.
Qo'llab-quvvatlanadigan formatlar:
    1. apple - olma
    apple – olma
    apple: olma
    apple = olma
    apple => olma
    apple    olma      (tab yoki 2+ probel)
    apple, olma
    CSV fayllar: (english, uzbek) yoki (english, uzbek, example)
    Docx jadvallari: har qatorda kamida 2 ustun [EN, UZ, EX]
"""
import re
import csv
from pathlib import Path

_DELIM_PATTERN = re.compile(r"\s*(?:->|=>|[-–—:=,\t]|\s{2,})\s*")
_LEADING_NUM_PATTERN = re.compile(r"^\s*(?:\[?\d+[\.\)\-\]]|[-*•])\s*")


def parse_line(line: str) -> tuple[str, str] | None:
    line = line.strip()
    if not line or line.startswith("#"):
        return None
    # Qator boshidagi raqamlash va markerlarni tozalash ("1. apple - olma" -> "apple - olma")
    line = _LEADING_NUM_PATTERN.sub("", line).strip()
    parts = _DELIM_PATTERN.split(line, maxsplit=1)
    if len(parts) != 2:
        return None
    eng, uz = parts[0].strip(), parts[1].strip()
    if not eng or not uz:
        return None
    return eng, uz


def parse_txt(path: str | Path) -> list[tuple]:
    pairs = []
    # Turli xil kodlashlarni xavfsiz o'qish (utf-8, utf-8-sig)
    with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
        for line in f:
            result = parse_line(line)
            if result:
                pairs.append(result)
    return pairs


def parse_csv(path: str | Path) -> list[tuple]:
    pairs = []
    with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
        # Delimiter aniqlash
        sample = f.read(2048)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            if not row or len(row) < 2:
                continue
            eng = _LEADING_NUM_PATTERN.sub("", row[0].strip()).strip()
            uz = row[1].strip()
            # Header qatorini tekshirish
            if eng.lower() in ("english", "en", "word", "so'z", "soz", "id") and uz.lower() in ("uzbek", "uz", "tarjima", "translation"):
                continue
            if not eng or not uz:
                continue
            # Agar 3-ustunda example bo'lsa
            ex = row[2].strip() if len(row) >= 3 else ""
            pairs.append((eng, uz, ex) if ex else (eng, uz))
    return pairs


def parse_docx(path: str | Path) -> list[tuple]:
    import docx  # python-docx

    doc = docx.Document(str(path))
    pairs = []

    # 1) Jadvallar (agar mavjud bo'lsa) — har qator: [EN, UZ, EX]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] and cells[1]:
                if cells[0].lower() in ("english", "en", "so'z", "word"):
                    continue  # header qatorini o'tkazib yuborish
                ex = cells[2].strip() if len(cells) >= 3 else ""
                pairs.append((cells[0], cells[1], ex) if ex else (cells[0], cells[1]))

    # 2) Oddiy paragraflar — matn qatorlari
    for para in doc.paragraphs:
        result = parse_line(para.text)
        if result:
            pairs.append(result)

    return pairs


def parse_file(path: str | Path) -> list[tuple]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return parse_txt(path)
    elif suffix == ".csv":
        return parse_csv(path)
    elif suffix == ".docx":
        return parse_docx(path)
    else:
        raise ValueError(f"Qo'llab-quvvatlanmaydigan fayl turi: {suffix} (.txt, .csv yoki .docx kerak)")
